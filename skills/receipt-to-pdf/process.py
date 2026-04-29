"""영수증 사진 → 자동 정리 → PPT 슬라이드 크기 PDF/PPTX

사용법:
    python process.py <input_dir> [output_name]
    python process.py <input_dir> [output_name] --points <points_dir>

기본 동작:
    - <input_dir>의 모든 .jpg/.jpeg/.png 파일을 영수증으로 인식
    - <input_dir>의 부모/<output_name 또는 폴더명>_rect/ 에 개별 정리 이미지 저장
    - <input_dir>의 부모/<output_name 또는 폴더명>.pdf, .pptx 생성

--points 옵션:
    - <points_dir>에 같은 파일명으로 빨간 점(4코너 마커) 찍힌 사본을 두면 그 좌표로 변환
"""
from __future__ import annotations
import argparse
import glob
import io
import os
import sys
from pathlib import Path

import numpy as np
from PIL import Image, ImageOps, ImageEnhance

# Windows cp949 회피
if sys.platform == "win32":
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
    sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding="utf-8", errors="replace")


# ─── 종이 마스크 / Otsu ────────────────────────────────────────
def paper_mask(arr: np.ndarray, bri: int = 165, sat: int = 70) -> np.ndarray:
    b = arr[:, :, :3].mean(2).astype(float)
    s = arr[:, :, :3].max(2).astype(float) - arr[:, :, :3].min(2).astype(float)
    return (b > bri) & (s < sat)


def otsu_threshold(arr: np.ndarray) -> int:
    g = arr[:, :, :3].mean(2).astype(float)
    hist, _ = np.histogram(g.flatten(), 256, (0, 256))
    total = g.size
    s_t = np.dot(np.arange(256), hist)
    sb, wb, mv, bt = 0, 0, 0, 0
    for t in range(256):
        wb += hist[t]
        if wb == 0:
            continue
        wf = total - wb
        if wf == 0:
            break
        sb += t * hist[t]
        v = wb * wf * (sb / wb - (s_t - sb) / wf) ** 2
        if v > mv:
            mv, bt = v, t
    return bt


# ─── 가장 큰 연결 영역만 남기기 (다른 흰 종이/배경 제거) ───────
def isolate_largest_blob(mask: np.ndarray, downscale: int = 8) -> np.ndarray:
    """4-연결 컴포넌트 중 가장 큰 것만 남긴다.
    속도를 위해 다운스케일된 마스크에서 iterative dilation으로 컴포넌트를 키우고,
    가장 큰 컴포넌트의 위치만 원본 해상도로 복원해 마스킹한다."""
    h, w = mask.shape
    sh, sw = h // downscale, w // downscale
    if sh < 4 or sw < 4:
        return mask

    # block-mean 다운샘플 + 50% 임계
    cropped = mask[: sh * downscale, : sw * downscale]
    small = cropped.reshape(sh, downscale, sw, downscale).mean(axis=(1, 3)) > 0.5
    if not small.any():
        return mask

    visited = np.zeros_like(small)
    best_size = 0
    best_comp = None

    while True:
        candidates = small & ~visited
        remaining = int(candidates.sum())
        # 남은 흰 픽셀이 현재 최선보다 적으면 더 큰 블롭 불가능 → 종료
        if remaining == 0 or remaining <= best_size:
            break

        ys, xs = np.where(candidates)
        comp = np.zeros_like(small)
        comp[ys[0], xs[0]] = True

        # iterative 4-연결 dilation
        prev_size = 0
        while True:
            d = comp.copy()
            d[1:, :] |= comp[:-1, :]
            d[:-1, :] |= comp[1:, :]
            d[:, 1:] |= comp[:, :-1]
            d[:, :-1] |= comp[:, 1:]
            comp = d & small
            sz = int(comp.sum())
            if sz == prev_size:
                break
            prev_size = sz

        if prev_size > best_size:
            best_size = prev_size
            best_comp = comp
        visited |= comp

    if best_comp is None:
        return mask

    # 원본 해상도로 복원
    big = best_comp.repeat(downscale, axis=0).repeat(downscale, axis=1)
    full = np.zeros_like(mask)
    full[: big.shape[0], : big.shape[1]] = big
    return mask & full


# ─── bbox + 엣지 직선 피팅 → 4코너 ────────────────────────────
def find_receipt_bbox(mask: np.ndarray, kernel: int = 80, ratio: float = 0.45):
    h, w = mask.shape
    rs = np.convolve(mask.mean(axis=1), np.ones(kernel) / kernel, mode="same")
    cs = np.convolve(mask.mean(axis=0), np.ones(kernel) / kernel, mode="same")
    rows = np.where(rs > rs.max() * ratio)[0]
    cols = np.where(cs > cs.max() * ratio)[0]
    if len(rows) == 0 or len(cols) == 0:
        return None
    pad = 100
    return (
        max(0, cols[0] - pad),
        max(0, rows[0] - pad),
        min(w, cols[-1] + pad),
        min(h, rows[-1] + pad),
    )


def fit_line_robust(pts: np.ndarray, axis: str = "x", n_iter: int = 4):
    if axis == "x":
        xs, ys = pts[:, 0].astype(float), pts[:, 1].astype(float)
    else:
        xs, ys = pts[:, 1].astype(float), pts[:, 0].astype(float)
    slope, ic = np.polyfit(xs, ys, 1)
    for _ in range(n_iter):
        res = np.abs(ys - (slope * xs + ic))
        med = max(np.median(res), 1.0)
        m = res < 3 * med + 5
        if m.sum() < 10:
            break
        xs, ys = xs[m], ys[m]
        slope, ic = np.polyfit(xs, ys, 1)
    return slope, ic


def find_corners_auto(mask: np.ndarray, min_span: int = 20):
    bbox = find_receipt_bbox(mask)
    if bbox is None:
        return None
    l, t, r, b = bbox
    m = mask.copy()
    m[:t, :] = False
    m[b:, :] = False
    m[:, :l] = False
    m[:, r:] = False

    h, w = m.shape
    top_pts, bot_pts = [], []
    for c in range(w):
        rows = np.where(m[:, c])[0]
        if len(rows) > min_span:
            top_pts.append((c, rows[0]))
            bot_pts.append((c, rows[-1]))
    left_pts, right_pts = [], []
    for r2 in range(h):
        cols = np.where(m[r2, :])[0]
        if len(cols) > min_span:
            left_pts.append((cols[0], r2))
            right_pts.append((cols[-1], r2))

    if min(len(top_pts), len(bot_pts), len(left_pts), len(right_pts)) < 30:
        return None

    top_pts, bot_pts = np.array(top_pts), np.array(bot_pts)
    left_pts, right_pts = np.array(left_pts), np.array(right_pts)

    def filt(pts, idx, asc=True, pct=20, slack=50):
        v = pts[:, idx]
        thr = np.percentile(v, pct if asc else 100 - pct)
        return pts[v < thr + slack] if asc else pts[v > thr - slack]

    ts, ti = fit_line_robust(filt(top_pts, 1, True), "x")
    bs, bi = fit_line_robust(filt(bot_pts, 1, False), "x")
    ls, li = fit_line_robust(filt(left_pts, 0, True), "y")
    rs, ri = fit_line_robust(filt(right_pts, 0, False), "y")

    def isect(hs, hi, vs, vi):
        x = (vs * hi + vi) / (1 - vs * hs)
        y = hs * x + hi
        return (x, y)

    return (
        isect(ts, ti, ls, li),
        isect(ts, ti, rs, ri),
        isect(bs, bi, rs, ri),
        isect(bs, bi, ls, li),
    )


# ─── 빨간 점 모드 ──────────────────────────────────────────────
def find_corners_from_red_dots(arr: np.ndarray):
    r = arr[:, :, 0].astype(float)
    g = arr[:, :, 1].astype(float)
    b = arr[:, :, 2].astype(float)
    red = (r > 180) & (r > g * 1.7) & (r > b * 1.7)
    if red.sum() == 0:
        red = (r > 150) & (r > g * 1.3) & (r > b * 1.3) & (g < 140) & (b < 140)

    rows, cols = np.where(red)
    if len(rows) == 0:
        return None
    pts = np.column_stack([cols, rows])

    used = np.zeros(len(pts), bool)
    clusters = []
    for i in range(len(pts)):
        if used[i]:
            continue
        d = np.linalg.norm(pts - pts[i], axis=1)
        near = d < 200
        clusters.append((int(near.sum()), pts[near].mean(0)))
        used[near] = True
    # 너무 큰 군집(피부 등) 제외
    small = [(sz, c) for sz, c in clusters if sz <= 500]
    if len(small) < 4:
        small = sorted(clusters, key=lambda x: x[0])[:4]
    if len(small) < 4:
        return None

    # 4개 점을 TL/TR/BR/BL로 분류
    pts4 = np.array([c for _, c in small[:4]], float)
    s, d = pts4[:, 0] + pts4[:, 1], pts4[:, 0] - pts4[:, 1]
    tl = pts4[s.argmin()]
    br = pts4[s.argmax()]
    tr = pts4[d.argmax()]
    bl = pts4[d.argmin()]
    return tuple(tl), tuple(tr), tuple(br), tuple(bl)


# ─── 원근 변환 + 화질 개선 ─────────────────────────────────────
def perspective_coeffs(src, dst):
    A, bv = [], []
    for (x, y), (X, Y) in zip(dst, src):
        A += [[x, y, 1, 0, 0, 0, -X * x, -X * y], [0, 0, 0, x, y, 1, -Y * x, -Y * y]]
        bv += [X, Y]
    return np.linalg.solve(np.array(A, float), np.array(bv, float)).tolist()


def warp(img: Image.Image, tl, tr, br, bl) -> Image.Image:
    pts = [np.array(c, float) for c in [tl, tr, br, bl]]
    tl, tr, br, bl = pts
    W = int(max(np.linalg.norm(tr - tl), np.linalg.norm(br - bl)))
    H = int(max(np.linalg.norm(bl - tl), np.linalg.norm(br - tr)))
    dst = np.array([[0, 0], [W - 1, 0], [W - 1, H - 1], [0, H - 1]], float)
    c = perspective_coeffs(np.array([tl, tr, br, bl], float), dst)
    return img.transform((W, H), Image.PERSPECTIVE, c, Image.BICUBIC)


def enhance(img: Image.Image) -> Image.Image:
    g = img.convert("L")
    return ImageEnhance.Contrast(ImageOps.autocontrast(g, cutoff=1)).enhance(1.4)


def auto_rotate(img: Image.Image) -> Image.Image:
    """영수증을 세로 방향으로 정렬한다.
    - 영수증은 본질적으로 세로가 긴 직사각형 → W>H면 90° 회전
    - 회전 방향(CW/CCW)은 헤더(상단)가 더 진한 쪽으로 결정 (휴리스틱 — 케이스에 따라 빗나갈 수 있음)
    - 180° 뒤집힘은 OCR 없이는 안정적으로 감지 불가 → --flip 옵션으로 수동 지정"""
    w, h = img.size
    if w <= h:
        return img

    a_ccw = np.array(img.rotate(90, expand=True).convert("L"))
    a_cw = np.array(img.rotate(-90, expand=True).convert("L"))
    thr = float(a_ccw.mean()) * 0.7
    top_ccw = int((a_ccw[: a_ccw.shape[0] // 3] < thr).sum())
    top_cw = int((a_cw[: a_cw.shape[0] // 3] < thr).sum())
    return img.rotate(90 if top_ccw >= top_cw else -90, expand=True)


# ─── 영수증 1개 처리 ───────────────────────────────────────────
def process_one(src_path: Path, points_dir: Path | None, flip: bool = False):
    img = Image.open(src_path)
    arr = np.array(img)

    corners = None
    if points_dir:
        pp = points_dir / src_path.name
        if pp.exists():
            corners = find_corners_from_red_dots(np.array(Image.open(pp)))

    if corners is None:
        otsu = otsu_threshold(arr)
        bri = max(otsu + 10, 150)
        mask = paper_mask(arr, bri=bri, sat=70)
        mask = isolate_largest_blob(mask)
        corners = find_corners_auto(mask)

    if corners is None:
        return None
    tl, tr, br, bl = corners
    out = auto_rotate(enhance(warp(img, tl, tr, br, bl)))
    if flip:
        out = out.rotate(180, expand=True)
    return out


# ─── PPT 슬라이드 합성 ─────────────────────────────────────────
SLIDE_W_IN = 13.333
SLIDE_H_IN = 7.5
DPI = 200
MARGIN_IN = 0.3
GAP_IN = 0.2
PER_ROW = 5


def compose_layout(rect_files: list[Path], output_pdf: Path, output_pptx: Path | None):
    slide_w_px = int(SLIDE_W_IN * DPI)
    slide_h_px = int(SLIDE_H_IN * DPI)
    margin_px = int(MARGIN_IN * DPI)
    gap_px = int(GAP_IN * DPI)
    slot_w_in = (SLIDE_W_IN - 2 * MARGIN_IN - (PER_ROW - 1) * GAP_IN) / PER_ROW
    slot_h_in = SLIDE_H_IN - 2 * MARGIN_IN
    slot_w_px = int(slot_w_in * DPI)
    slot_h_px = int(slot_h_in * DPI)

    def fit(im: Image.Image, sw: int, sh: int) -> Image.Image:
        w, h = im.size
        sc = min(sw / w, sh / h)
        return im.resize((int(w * sc), int(h * sc)), Image.LANCZOS)

    imgs = [Image.open(f) for f in rect_files]

    # PDF (PIL 합성, 좌상단 정렬)
    slides = []
    for s in range(0, len(imgs), PER_ROW):
        chunk = imgs[s:s + PER_ROW]
        fitted = [fit(im, slot_w_px, slot_h_px) for im in chunk]
        canvas = Image.new("RGB", (slide_w_px, slide_h_px), "white")
        x = margin_px
        y = margin_px
        for im in fitted:
            canvas.paste(im, (x, y))
            x += slot_w_px + gap_px
        slides.append(canvas)

    slides[0].save(
        output_pdf,
        save_all=True,
        append_images=slides[1:],
        resolution=DPI,
    )

    # PPTX (선택)
    if output_pptx is not None:
        try:
            from pptx import Presentation
            from pptx.util import Inches
        except ImportError:
            print("⚠ python-pptx 없음 — PPTX 생략")
            return

        prs = Presentation()
        prs.slide_width = Inches(SLIDE_W_IN)
        prs.slide_height = Inches(SLIDE_H_IN)
        blank = prs.slide_layouts[6]

        for s in range(0, len(imgs), PER_ROW):
            slide = prs.slides.add_slide(blank)
            chunk_files = rect_files[s:s + PER_ROW]
            chunk = imgs[s:s + PER_ROW]
            x = MARGIN_IN
            y = MARGIN_IN
            for img_path, im in zip(chunk_files, chunk):
                w, h = im.size
                sc = min(slot_w_in / (w / DPI), slot_h_in / (h / DPI))
                fw, fh = w * sc / DPI, h * sc / DPI
                slide.shapes.add_picture(
                    str(img_path), Inches(x), Inches(y),
                    width=Inches(fw), height=Inches(fh),
                )
                x += slot_w_in + GAP_IN

        prs.save(output_pptx)


# ─── 메인 ──────────────────────────────────────────────────────
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("input_dir", type=Path)
    ap.add_argument("output_name", nargs="?", default=None,
                    help="출력 파일 이름 (기본: input_dir 폴더명)")
    ap.add_argument("--points", type=Path, default=None,
                    help="빨간 점 마커 이미지 폴더")
    ap.add_argument("--flip", default="",
                    help="180° 뒤집을 파일 이름(쉼표 구분, stem 일치, 부분 매칭). "
                         "예: --flip 094610,094719")
    ap.add_argument("--no-pptx", action="store_true")
    args = ap.parse_args()
    flip_keys = [k.strip() for k in args.flip.split(",") if k.strip()]

    inp = args.input_dir.resolve()
    if not inp.is_dir():
        print(f"❌ 폴더 없음: {inp}")
        sys.exit(1)

    base = args.output_name or inp.name
    parent = inp.parent
    rect_dir = parent / f"{base}_rect"
    rect_dir.mkdir(exist_ok=True)

    pdf_path = parent / f"{base}.pdf"
    pptx_path = None if args.no_pptx else parent / f"{base}.pptx"

    files = []
    for ext in ("*.jpg", "*.jpeg", "*.png", "*.JPG", "*.JPEG", "*.PNG"):
        files.extend(inp.glob(ext))
    files = sorted(set(files))

    if not files:
        print(f"❌ 이미지 없음: {inp}")
        sys.exit(1)

    print(f"입력: {inp}  ({len(files)}개)")
    print(f"개별 출력: {rect_dir}")
    print(f"PDF: {pdf_path}")
    if pptx_path:
        print(f"PPTX: {pptx_path}")
    print()

    rect_files = []
    for fp in files:
        print(f"▶ {fp.name}")
        do_flip = any(k in fp.stem for k in flip_keys)
        result = process_one(fp, args.points, flip=do_flip)
        if result is None:
            print("  ⚠ 코너 감지 실패 — 건너뜀")
            continue
        out = rect_dir / f"{fp.stem}_rect.jpg"
        result.save(out, quality=95)
        flip_tag = " [flipped]" if do_flip else ""
        print(f"  ✓ {out.name} ({result.size[0]}x{result.size[1]}){flip_tag}")
        rect_files.append(out)

    if not rect_files:
        print("\n❌ 처리된 영수증 없음")
        sys.exit(1)

    print(f"\n레이아웃 생성 중... ({len(rect_files)}개 → {(len(rect_files)+PER_ROW-1)//PER_ROW}슬라이드)")
    compose_layout(rect_files, pdf_path, pptx_path)
    print(f"✓ 완료")


if __name__ == "__main__":
    main()
